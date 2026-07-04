#!/usr/bin/env python3
"""
dataconnect_export.py — one-shot xlsx -> per-class JSON export for the DataConnect shim demo.

Reads the enriched demo workbook (V5_Demo_Package_Enriched.xlsx) and writes one JSON array
per sheet/"class" to tools/dataconnect-data/<class>.json, so the shim (dataconnect_shim.py)
and eventually the real DataConnect instance can be swapped in behind the same client with
no code change (see docs/superpowers/specs/2026-07-04-dataconnect-scenario-a-demo-design.md).

Sheets read (class -> output slug):
  Asset Registry           -> asset_registry.json
  Work Orders              -> work_orders.json
  Tasks                    -> tasks.json
  Tickets                  -> tickets.json
  Safety_Inspections_V3    -> safety_inspections_v3.json
  Roadway_Inspections_V3   -> roadway_inspections_v3.json
  ITS_Inspections_V3       -> its_inspections_v3.json
  Incidents_V3             -> incidents_v3.json

Sheets skipped (empty in the source workbook, header row only): Pavement, Condtion History (sic).

Coordinate remediation (Asset Registry / Tasks / Tickets / Roadway_Inspections_V3 /
ITS_Inspections_V3 — the sheets that carry lon/lat columns). Three repairs are applied,
in this order of preference:
  - already inside the I-595 corridor bounds (lon -80.5..-80.0, lat 25.9..26.3)  -> kept as-is.
  1. missing decimal point: a coordinate with a dropped decimal point, e.g. -80329552,
     is repaired by a magnitude heuristic — both corridor coordinates have a 2-digit
     integer part (80.x / 26.x), so the point is reinserted after the first 2 digits
     -> -80.329552.
  2. sign-flipped longitude: +80.32 where the corridor is west (-80.32) -> negated.
  3. lon/lat axis-swap: if the x column holds a latitude and the y column holds a longitude
     (after applying repairs 1-2 to each candidate value), the pair is swapped back. This is
     deterministic on this corridor: the lat band (25.9..26.3) and the lon band (-80.5..-80.0)
     cannot overlap, so a swapped pair can never be mistaken for an in-bounds one. Axis-swap
     repairs are additionally tagged `_repair: "axis-swap"` on the record. In this workbook
     ITS_Inspections_V3's x/y columns are swapped for the *entire* sheet, and ~30% of Asset
     Registry rows are swapped per-row.
  Repaired rows count as repaired (kept), carry `_coordinateRemediation` with the reason,
  and anything still out of bounds after all three repairs is quarantined with a reason.
  Counts of ok/repaired/quarantined are printed per class.

  Coordinates are only load-bearing for Asset Registry (they place the pins in the map
  layer), so an Asset Registry row with unrepairable coordinates is dropped from
  asset_registry.json entirely and moved to asset_registry.rejected.json. On the other
  coordinate-bearing sheets (Tasks/Tickets/Roadway_Inspections_V3/ITS_Inspections_V3) the
  coordinates are incidental context — those rows are joined into scoring by Asset ID, not
  placed independently — so an unrepairable coordinate there just gets nulled out in place
  (tagged `_coordinateIssue`) and copied into <slug>.rejected.json for visibility, but the
  row itself (and its non-coordinate fields) is kept.

Usage:
  python3 tools/dataconnect_export.py <path/to/V5_Demo_Package_Enriched.xlsx>
"""
from __future__ import annotations

import datetime
import json
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(HERE, "dataconnect-data")

# I-595 corridor bounds — matches the design spec §Components.1.
LON_MIN, LON_MAX = -80.5, -80.0
LAT_MIN, LAT_MAX = 25.9, 26.3

# sheet name -> (output slug, (lon_column, lat_column) or None, drop_row_on_bad_coords)
# drop_row_on_bad_coords=True only for Asset Registry: its coordinates are load-bearing (map
# pin placement in the Cesium asset layer). On the other sheets coordinates are incidental
# context re-joined by Asset ID, so a bad coordinate there nulls the two fields in place
# rather than discarding an otherwise-good row.
SHEETS = {
    "Asset Registry":          ("asset_registry",          ("X Coordinates", "Y Coordinates"), True),
    "Pavement":                None,   # empty in the source workbook — skipped
    "Condtion History":        None,   # empty in the source workbook (sic) — skipped
    "Work Orders":             ("work_orders",             None,                               False),
    "Tasks":                   ("tasks",                   ("X Coordinate", "Y Coordinate"),   False),
    "Tickets":                 ("tickets",                 ("X Coordinate", "Y Coordinate"),   False),
    "Safety_Inspections_V3":   ("safety_inspections_v3",   None,                               False),
    "Roadway_Inspections_V3":  ("roadway_inspections_v3",  ("x_coordinate", "y_coordinate"),   False),
    "ITS_Inspections_V3":      ("its_inspections_v3",      ("x_coordinates", "y_coordinates"), False),
    "Incidents_V3":            ("incidents_v3",            None,                               False),
}


def _dedupe_headers(header: tuple) -> list[str]:
    """Strip whitespace and disambiguate duplicate column names (Asset Registry has two
    columns literally both named 'Segment') so building a dict from the row never silently
    drops a column."""
    seen: dict[str, int] = {}
    out = []
    for raw in header:
        name = str(raw).strip() if raw is not None else ""
        seen[name] = seen.get(name, 0) + 1
        out.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return out


def _json_safe(value):
    """Convert an openpyxl cell value into something json.dump can serialize."""
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and (value != value):  # NaN
        return None
    return value


def _to_float(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip())
        except ValueError:
            return None
    return None


def _in_bounds(lon, lat) -> bool:
    return (
        lon is not None and lat is not None
        and LON_MIN <= lon <= LON_MAX
        and LAT_MIN <= lat <= LAT_MAX
    )


def _fix_missing_decimal(value):
    """Magnitude heuristic: a coordinate with its decimal point dropped reads as a large
    whole number (e.g. -80329552). Both corridor coordinates have a 2-digit integer part
    (80.x for longitude, 26.x for latitude), so reinsert the point after the first 2 digits.
    Returns None if `value` doesn't look like this defect."""
    if value is None or not float(value).is_integer() or abs(value) < 1000:
        return None
    digits = str(int(value))
    neg = digits.startswith("-")
    digits = digits[1:] if neg else digits
    if len(digits) < 3:
        return None
    fixed = float(f"{digits[:2]}.{digits[2:]}")
    return -fixed if neg else fixed


def _lon_candidates(value):
    """Yield (candidate, repair_description or None) longitude interpretations of a raw
    value, in preference order: as-is, missing-decimal repaired, sign-flip negated, both."""
    yield value, None
    fixed = _fix_missing_decimal(value)
    if fixed is not None:
        yield fixed, "longitude missing decimal point (magnitude heuristic)"
    if value > 0:
        yield -value, "sign-flipped longitude negated"
    if fixed is not None and fixed > 0:
        yield -fixed, "longitude missing decimal point + sign-flipped"


def _lat_candidates(value):
    """Yield (candidate, repair_description or None) latitude interpretations of a raw
    value. Latitudes here are always positive (26.x), so no sign-flip variant."""
    yield value, None
    fixed = _fix_missing_decimal(value)
    if fixed is not None:
        yield fixed, "latitude missing decimal point (magnitude heuristic)"


def remediate_coords(lon_raw, lat_raw):
    """Returns (lon, lat, status, repair_kind, reason) where status is one of "ok" /
    "repaired" / "quarantined", repair_kind is "axis-swap" for swapped-column repairs
    (else None), and `reason` is None for "ok"."""
    lon = _to_float(lon_raw)
    lat = _to_float(lat_raw)
    if lon is None or lat is None:
        return None, None, "quarantined", None, "non-numeric coordinate value"

    # Direct orientation: x column is longitude, y column is latitude.
    for cand_lon, lon_repair in _lon_candidates(lon):
        for cand_lat, lat_repair in _lat_candidates(lat):
            if _in_bounds(cand_lon, cand_lat):
                if lon_repair is None and lat_repair is None:
                    return cand_lon, cand_lat, "ok", None, None
                reason = " + ".join(r for r in (lon_repair, lat_repair) if r)
                return cand_lon, cand_lat, "repaired", None, reason

    # Axis-swap orientation: x column holds the latitude, y column the longitude.
    # Deterministic on this corridor — the lat band (25.9..26.3) and lon band
    # (-80.5..-80.0) cannot overlap, so a swapped pair is unambiguous.
    for cand_lon, lon_repair in _lon_candidates(lat):
        for cand_lat, lat_repair in _lat_candidates(lon):
            if _in_bounds(cand_lon, cand_lat):
                reason = " + ".join(
                    r for r in ("lon/lat axis-swap corrected", lon_repair, lat_repair) if r
                )
                return cand_lon, cand_lat, "repaired", "axis-swap", reason

    return lon, lat, "quarantined", None, f"out of corridor bounds (lon={lon}, lat={lat})"


def export_sheet(
    wb, sheet_name: str, slug: str, coord_cols, drop_row_on_bad_coords: bool
) -> tuple[int, int, int]:
    """Writes <slug>.json (+ <slug>.rejected.json if anything was quarantined).
    Returns (kept_count, repaired_count, quarantined_count) — quarantined_count is rows
    dropped from the main file (drop_row_on_bad_coords=True) or rows kept with their
    coordinates nulled out (drop_row_on_bad_coords=False); both land in rejected.json."""
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        print(f"[dataconnect_export] {sheet_name}: no header row, skipping")
        return 0, 0, 0
    keys = _dedupe_headers(header)

    lon_idx = lat_idx = None
    if coord_cols is not None:
        lon_col, lat_col = coord_cols
        if lon_col in keys and lat_col in keys:
            lon_idx, lat_idx = keys.index(lon_col), keys.index(lat_col)
        else:
            print(f"[dataconnect_export] WARNING: {sheet_name} missing expected coordinate "
                  f"columns {lon_col!r}/{lat_col!r}; skipping coordinate remediation")

    kept, repaired, quarantined = [], 0, []
    for row in rows:
        if row is None or all(v is None for v in row):
            continue
        record = {k: _json_safe(v) for k, v in zip(keys, row)}

        if lon_idx is not None:
            lon_col, lat_col = coord_cols
            lon, lat, status, repair_kind, reason = remediate_coords(row[lon_idx], row[lat_idx])
            if status == "quarantined":
                if drop_row_on_bad_coords:
                    record["_rejectReason"] = reason
                    quarantined.append(record)
                    continue
                # Coordinates are incidental on this sheet (joined by Asset ID, not placed
                # independently) — null them out but keep the row and its other fields.
                record[lon_col] = None
                record[lat_col] = None
                record["_coordinateIssue"] = reason
                quarantined.append({**record, "_rejectReason": reason})
            elif status == "repaired":
                record[lon_col] = lon
                record[lat_col] = lat
                record["_coordinateRemediation"] = reason
                if repair_kind == "axis-swap":
                    record["_repair"] = "axis-swap"
                repaired += 1

        kept.append(record)

    os.makedirs(OUT_DIR, exist_ok=True)
    out_path = os.path.join(OUT_DIR, f"{slug}.json")
    with open(out_path, "w") as f:
        json.dump(kept, f, indent=2)
        f.write("\n")

    if quarantined:
        rej_path = os.path.join(OUT_DIR, f"{slug}.rejected.json")
        with open(rej_path, "w") as f:
            json.dump(quarantined, f, indent=2)
            f.write("\n")

    total = len(kept) + (len(quarantined) if drop_row_on_bad_coords else 0)
    verb = "dropped" if drop_row_on_bad_coords else "coordinates nulled (row kept)"
    print(f"[dataconnect_export] {sheet_name} -> {slug}.json: "
          f"{total} rows read, {len(kept)} kept ({repaired} repaired), "
          f"{len(quarantined)} quarantined ({verb})"
          + (f" -> {slug}.rejected.json" if quarantined else ""))
    return len(kept), repaired, len(quarantined)


def main():
    if len(sys.argv) < 2:
        sys.exit("Usage: python3 tools/dataconnect_export.py <path/to/V5_Demo_Package_Enriched.xlsx>")
    xlsx_path = sys.argv[1]
    if not os.path.exists(xlsx_path):
        sys.exit(f"[dataconnect_export] file not found: {xlsx_path}")

    print(f"[dataconnect_export] loading {xlsx_path} ...")
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    totals = {"kept": 0, "repaired": 0, "quarantined": 0}
    for sheet_name, spec in SHEETS.items():
        if sheet_name not in wb.sheetnames:
            print(f"[dataconnect_export] WARNING: sheet {sheet_name!r} not found in workbook, skipping")
            continue
        if spec is None:
            print(f"[dataconnect_export] {sheet_name}: skipped (known-empty sheet)")
            continue
        slug, coord_cols, drop_row_on_bad_coords = spec
        kept, repaired, quarantined = export_sheet(
            wb, sheet_name, slug, coord_cols, drop_row_on_bad_coords
        )
        totals["kept"] += kept
        totals["repaired"] += repaired
        totals["quarantined"] += quarantined

    print(f"[dataconnect_export] done. total kept={totals['kept']} "
          f"(repaired={totals['repaired']}), quarantined={totals['quarantined']}")
    print(f"[dataconnect_export] output: {OUT_DIR}")


if __name__ == "__main__":
    main()
